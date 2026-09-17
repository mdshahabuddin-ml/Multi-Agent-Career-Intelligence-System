import io
import logging
from typing import Optional

logger = logging.getLogger(__name__)


def _extract_pdf_text(pdf_reader) -> str:
    """Extract text across all pages, tolerating per-page failures."""
    text = ""
    for page in pdf_reader.pages:
        try:
            page_text = page.extract_text()
        except Exception as e:
            logger.warning(f"Skipping unreadable PDF page: {e}")
            continue
        if page_text:
            text += page_text + "\n"
    return text.strip()


def parse_pdf(file_content: bytes) -> str:
    """Extract text from PDF file."""
    try:
        import pypdf
        pdf_reader = pypdf.PdfReader(io.BytesIO(file_content))
        text = _extract_pdf_text(pdf_reader)
        if not text:
            # Fallback: layout mode recovers text from PDFs whose default
            # extraction yields nothing (custom encodings / odd content streams).
            try:
                layout_text = ""
                for page in pdf_reader.pages:
                    try:
                        page_text = page.extract_text(extraction_mode="layout")
                    except Exception as e:
                        logger.warning(f"Skipping unreadable PDF page (layout mode): {e}")
                        continue
                    if page_text:
                        layout_text += page_text + "\n"
                text = layout_text.strip()
            except Exception as e:
                logger.warning(f"PDF layout-mode extraction failed: {e}")
        if not text:
            raise ValueError(
                "No extractable text found in PDF. "
                "If this is a scanned document, please upload a text-based PDF."
            )
        return text
    except ValueError:
        raise
    except Exception as e:
        logger.error(f"PDF parsing failed: {e}")
        raise ValueError(f"Failed to parse PDF: {e}")


def parse_docx(file_content: bytes) -> str:
    """Extract text from DOCX file."""
    # Legacy .doc (OLE compound document) is not readable by python-docx.
    # Detect it by magic bytes and fail with actionable guidance instead
    # of a cryptic traceback.
    if file_content[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        raise ValueError(
            "Legacy .doc format is not supported. "
            "Please re-save the file as .docx and upload again."
        )
    try:
        import docx
        doc = docx.Document(io.BytesIO(file_content))
        text = ""
        for para in doc.paragraphs:
            if para.text:
                text += para.text + "\n"
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text:
                        text += cell.text + "\n"
        return text.strip()
    except Exception as e:
        logger.error(f"DOCX parsing failed: {e}")
        raise ValueError(f"Failed to parse DOCX: {e}")


def parse_text(file_content: bytes) -> str:
    """Extract text from plain text file."""
    try:
        return file_content.decode("utf-8").strip()
    except UnicodeDecodeError:
        try:
            return file_content.decode("latin-1").strip()
        except Exception as e:
            logger.error(f"Text parsing failed: {e}")
            raise ValueError(f"Failed to parse text file: {e}")


def parse_document(file_content: bytes, mime_type: str, filename: str) -> str:
    """Parse document based on mime type or file extension."""
    mime_type = mime_type.lower()
    filename = filename.lower()

    if mime_type == "application/pdf" or filename.endswith(".pdf"):
        return parse_pdf(file_content)
    elif mime_type in [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/msword",
    ] or filename.endswith((".docx", ".doc")):
        return parse_docx(file_content)
    elif mime_type == "text/plain" or filename.endswith(".txt"):
        return parse_text(file_content)
    else:
        raise ValueError(f"Unsupported file type: {mime_type} ({filename})")


def extract_sections(text: str) -> dict:
    """Extract common resume sections from text."""
    import re

    sections = {
        "summary": "",
        "experience": "",
        "education": "",
        "skills": "",
        "projects": "",
        "certifications": "",
        "other": "",
    }

    section_patterns = {
        "summary": [
            r"(?:professional\s+)?summary",
            r"objective",
            r"profile",
            r"about\s+me",
        ],
        "experience": [
            r"experience",
            r"work\s+history",
            r"employment",
            r"professional\s+experience",
        ],
        "education": [
            r"education",
            r"academic\s+background",
            r"qualifications",
        ],
        "skills": [
            r"skills",
            r"technical\s+skills",
            r"competencies",
            r"technologies",
        ],
        "projects": [
            r"projects",
            r"personal\s+projects",
            r"key\s+projects",
        ],
        "certifications": [
            r"certifications",
            r"certificates",
            r"licenses",
        ],
    }

    lines = text.split("\n")
    current_section = "other"
    section_content = {k: [] for k in sections}

    for line in lines:
        line_lower = line.lower().strip()
        matched = False

        for section, patterns in section_patterns.items():
            for pattern in patterns:
                if re.match(rf"^{pattern}[:\s]*$", line_lower):
                    current_section = section
                    matched = True
                    break
            if matched:
                break

        if not matched and line.strip():
            section_content[current_section].append(line)

    for section, content in section_content.items():
        sections[section] = "\n".join(content).strip()

    return sections